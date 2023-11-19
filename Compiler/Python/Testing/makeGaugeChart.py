from openpyxl import Workbook

from openpyxl.chart import (
    DoughnutChart,
    Reference,
    Series,
)
from openpyxl.chart.series import DataPoint

theta = 240
total = 360
percent = .3

print(400 % 360)

ang = (( 0 - theta/2 + theta * percent -1) + 360) % 360
white = total - theta
bg1 = round((theta - 2) * percent, 3) 
bg2 = round((theta - 2) * (1- percent), 3)


data = [
    ["Total", total, total],
    ['Needle', 2, 2],
    ["bg_two", bg2, bg2],
    ['White', white, white],
    ["bg_one", bg1, bg1],
    [],
    ["Total", total, total],
    ['Needle', 2, 2],
    ["bg_two", total - 2, total - 2]
]

print(data)

wb = Workbook()
ws = wb.active

for row in data:
    ws.append(row)

chart = DoughnutChart(firstSliceAng=ang)
chart.title = "Gauge Chart"
chart.legend = None

data_two = Reference(ws, min_col=2, min_row=7, max_row=9)
chart.add_data(data_two, titles_from_data=True)
slices_two = [DataPoint(idx=i) for i in range(2)]
needle_inner, white_inner = slices_two
chart.series[0].data_points = slices_two
white_inner.graphicalProperties.solidFill = "FFFFFF"#"441bbf"
needle_inner.graphicalProperties.solidFill = "000000"
data_two = Reference(ws, min_col=3, min_row=1, max_row=3)


data = Reference(ws, min_col=2, min_row=1, max_row=5)
chart.add_data(data, titles_from_data=True)
slices = [DataPoint(idx=i) for i in range(4)]
needle, bg_two, white, bg_one = slices
chart.series[1].data_points = slices
white.graphicalProperties.solidFill = "FFFFFF"#"441bbf"
needle.graphicalProperties.solidFill = "000000"
bg_two.graphicalProperties.solidFill = "f54842"
bg_one.graphicalProperties.solidFill = "f54842"

chart.width = 10




ws.add_chart(chart, "E17")

wb.save("(Pre)ScoutingDataCompiler/Compiler/Output/gaugeTesting.xlsx")





# from openpyxl import Workbook
# from openpyxl.chart import PieChart, DoughnutChart, Series, Reference
# from openpyxl.chart.series import DataPoint


# data = [
#     ["Donut", "Pie"],
#     [25, 75],
#     [50, 1],
#     [25, 124],
#     [100],
# ]

# # based on http://www.excel-easy.com/examples/gauge-chart.html

# wb = Workbook()
# ws = wb.active
# for row in data:
#     ws.append(row)


# class GaugeChart:
#     from openpyxl import Workbook
#     from openpyxl.chart import PieChart, DoughnutChart, Series, Reference
#     def __init__(self, workbook, value: float, theta: int = 180, chart_range = range(0, 101, 1), title = "Gauge Chart") -> None:
#         self.workbook = workbook
#         self.doughnut_chart = DoughnutChart(title=title)
#         self.pie_slice = PieChart()
        
        
    
#     def get(self): return self.doughnut_chart + self.pie_slice
    


    
# # First chart is a doughnut chart
# c1 = DoughnutChart(firstSliceAng=270, holeSize=50)
# c1.title = "Code coverage"
# c1.legend = None

# ref = Reference(ws, min_col=1, min_row=2, max_row=5)
# s1 = Series(ref, title_from_data=False)

# slices = [DataPoint(idx=i) for i in range(4)]
# slices[0].graphicalProperties.solidFill = "FF3300" # red
# slices[1].graphicalProperties.solidFill = "FCF305" # yellow
# slices[2].graphicalProperties.solidFill = "1FB714" # green
# slices[3].graphicalProperties.noFill = True # invisible

# s1.data_points = slices
# c1.series = [s1]

# # Second chart is a pie chart
# c2 = PieChart(firstSliceAng=270)
# c2.legend = None

# ref = Reference(ws, min_col=2, min_row=2, max_col=2, max_row=4)
# s2 = Series(ref, title_from_data=False)

# slices = [DataPoint(idx=i) for i in range(3)]
# slices[0].graphicalProperties.noFill = True # invisible
# slices[1].graphicalProperties.solidFill = "000000" # black needle
# slices[2].graphicalProperties.noFill = True # invisible
# s2.data_points = slices
# c2.series = [s2]

# c1 += c2 # combine charts

# ws.add_chart(c1, "D1")

# wb.save("(Pre)ScoutingDataCompiler/Compiler/Output/gaugeTesting.xlsx")